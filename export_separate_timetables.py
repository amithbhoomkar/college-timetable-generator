import os
import json
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
from copy import copy

EXCEL_FILE = "Complete_Timetable.xlsx"
EXPORT_DIR = "Timetable_Exports"

def export_separate():
    print("Reading Complete_Timetable.xlsx...")
    wb = openpyxl.load_workbook(EXCEL_FILE)

    os.makedirs(f"{EXPORT_DIR}/Sections", exist_ok=True)
    os.makedirs(f"{EXPORT_DIR}/Departments", exist_ok=True)
    os.makedirs(f"{EXPORT_DIR}/Staff", exist_ok=True)
    os.makedirs(f"{EXPORT_DIR}/Rooms", exist_ok=True)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    bold_font = Font(name="Calibri", bold=True, size=10)
    regular_font = Font(name="Calibri", size=9)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
    )

    def apply_1page_setup(ws):
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 1. Export each Section to its own separate Excel file
    section_sheets = [s for s in wb.sheetnames if s.startswith("ISE_S") or s.startswith("CSBS_S")]
    print(f"Exporting {len(section_sheets)} section workbooks...")
    for s_name in section_sheets:
        src_ws = wb[s_name]
        new_wb = Workbook()
        dest_ws = new_wb.active
        dest_ws.title = s_name
        apply_1page_setup(dest_ws)

        for r in range(1, src_ws.max_row + 1):
            dest_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
            for c in range(1, src_ws.max_column + 1):
                src_cell = src_ws.cell(row=r, column=c)
                dest_cell = dest_ws.cell(row=r, column=c, value=src_cell.value)
                if src_cell.has_style:
                    dest_cell.font = copy(src_cell.font)
                    dest_cell.border = copy(src_cell.border)
                    dest_cell.fill = copy(src_cell.fill)
                    dest_cell.number_format = copy(src_cell.number_format)
                    dest_cell.protection = copy(src_cell.protection)
                    dest_cell.alignment = copy(src_cell.alignment)

        dest_ws.column_dimensions["A"].width = 15
        for col in range(2, 8):
            dest_ws.column_dimensions[get_column_letter(col)].width = 19

        new_wb.save(f"{EXPORT_DIR}/Sections/{s_name}.xlsx")

    # 2. Export Department Masters
    dept_sheets = ["ISE_MASTER", "CSBS_MASTER"]
    for d_name in dept_sheets:
        src_ws = wb[d_name]
        new_wb = Workbook()
        dest_ws = new_wb.active
        dest_ws.title = d_name
        apply_1page_setup(dest_ws)

        for r in range(1, src_ws.max_row + 1):
            dest_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height or 20
            for c in range(1, src_ws.max_column + 1):
                src_cell = src_ws.cell(row=r, column=c)
                dest_cell = dest_ws.cell(row=r, column=c, value=src_cell.value)
                if src_cell.has_style:
                    dest_cell.font = copy(src_cell.font)
                    dest_cell.border = copy(src_cell.border)
                    dest_cell.fill = copy(src_cell.fill)
                    dest_cell.alignment = copy(src_cell.alignment)

        for col in range(1, 8):
            dest_ws.column_dimensions[get_column_letter(col)].width = 18

        new_wb.save(f"{EXPORT_DIR}/Departments/{d_name}.xlsx")

    # 3. Export each Staff member to their own separate timetable
    staff_ws = wb["STAFF_TIMETABLE"]
    staff_events = defaultdict(list)
    for r in range(2, staff_ws.max_row + 1):
        row_vals = [staff_ws.cell(row=r, column=c).value for c in range(1, 9)]
        teacher_name = row_vals[0]
        staff_events[teacher_name].append(row_vals)

    print(f"Exporting {len(staff_events)} individual faculty workbooks...")
    for t_name, events in staff_events.items():
        safe_name = t_name.replace(" ", "_")
        new_wb = Workbook()
        dest_ws = new_wb.active
        dest_ws.title = safe_name[:31]
        apply_1page_setup(dest_ws)

        dest_ws.merge_cells("A1:F1")
        dest_ws["A1"] = f"FACULTY TIMETABLE - {t_name}"
        dest_ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="1F4E78")
        dest_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        dest_ws.row_dimensions[1].height = 25

        dest_ws.merge_cells("A2:F2")
        dest_ws["A2"] = f"Designation: {events[0][1]} | Total Assigned Load: {sum(3 if e[7]=='LAB' else 1 for e in events)} Hours/Week"
        dest_ws["A2"].font = Font(name="Calibri", bold=True, size=10, italic=True, color="595959")
        dest_ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        dest_ws.row_dimensions[2].height = 18

        headers = ["Day", "Time", "Subject", "Section", "Room", "Type"]
        dest_ws.row_dimensions[4].height = 22
        for c, h in enumerate(headers, 1):
            cell = dest_ws.cell(row=4, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = thin_border

        for r_idx, ev in enumerate(events, 5):
            vals = [ev[2], ev[3], ev[4], ev[5], ev[6], ev[7]]
            dest_ws.row_dimensions[r_idx].height = 19
            for c_idx, val in enumerate(vals, 1):
                cell = dest_ws.cell(row=r_idx, column=c_idx, value=val)
                cell.alignment = center
                cell.border = thin_border
                cell.font = regular_font

        for col in range(1, 7):
            dest_ws.column_dimensions[get_column_letter(col)].width = 20

        new_wb.save(f"{EXPORT_DIR}/Staff/{safe_name}.xlsx")

    # 4. Export each Room to its own separate timetable
    room_ws = wb["ROOM_UTILIZATION"]
    room_events = defaultdict(list)
    for r in range(2, room_ws.max_row + 1):
        row_vals = [room_ws.cell(row=r, column=c).value for c in range(1, 8)]
        room_name = row_vals[0]
        room_events[room_name].append(row_vals)

    print(f"Exporting {len(room_events)} individual room workbooks...")
    for r_name, events in room_events.items():
        new_wb = Workbook()
        dest_ws = new_wb.active
        dest_ws.title = r_name
        apply_1page_setup(dest_ws)

        dest_ws.merge_cells("A1:F1")
        dest_ws["A1"] = f"ROOM UTILIZATION TIMETABLE - {r_name}"
        dest_ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="1F4E78")
        dest_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        dest_ws.row_dimensions[1].height = 25

        headers = ["Day", "Time", "Subject", "Section", "Teacher", "Type"]
        dest_ws.row_dimensions[3].height = 22
        for c, h in enumerate(headers, 1):
            cell = dest_ws.cell(row=3, column=c, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = thin_border

        for r_idx, ev in enumerate(events, 4):
            vals = [ev[1], ev[2], ev[3], ev[4], ev[5], ev[6]]
            dest_ws.row_dimensions[r_idx].height = 19
            for c_idx, val in enumerate(vals, 1):
                cell = dest_ws.cell(row=r_idx, column=c_idx, value=val)
                cell.alignment = center
                cell.border = thin_border
                cell.font = regular_font

        for col in range(1, 7):
            dest_ws.column_dimensions[get_column_letter(col)].width = 20

        new_wb.save(f"{EXPORT_DIR}/Rooms/{r_name}.xlsx")

    print("\nAll separate timetable workbooks successfully created with 1-page fit setup!")

if __name__ == "__main__":
    export_separate()
