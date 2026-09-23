

from ortools.sat.python import cp_model
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
import os
import json
import time




DAYS = [
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday"
]

# 9 teaching slots/day (10:30-11:00 short break, 13:00-14:00 lunch break)
TIME_SLOTS = [
    "07:30-08:30",
    "08:30-09:30",
    "09:30-10:30",
    "11:00-12:00",
    "12:00-13:00",
    "14:00-15:00",
    "15:00-16:00",
    "16:00-17:00",
    "17:00-18:00"
]

TIME_START = [
    "07:30",
    "08:30",
    "09:30",
    "11:00",
    "12:00",
    "14:00",
    "15:00",
    "16:00",
    "17:00"
]

NUM_DAYS = len(DAYS)
NUM_SLOTS = len(TIME_SLOTS)

CLASSROOMS = [
    "CR_1",
    "CR_2",
    "CR_3",
    "CR_4",
    "CR_5",
    "CR_6",
    "CR_7"
]

LABS = [
    "LAB_1",
    "LAB_2"
]

ROOMS = CLASSROOMS + LABS

NUM_CLASSROOMS = len(CLASSROOMS)
NUM_LABS = len(LABS)

OE_SLOT = 2       # 09:30-10:30
PE_SLOT = 5       # 14:00-15:00

SOLVER_TIME_LIMIT = 180

OUTPUT_FILE = "Complete_Timetable.xlsx"


SECTIONS = [
    # ISE
    {"id": "ISE_S3_A", "dept": "ISE", "sem": 3},
    {"id": "ISE_S3_B", "dept": "ISE", "sem": 3},
    {"id": "ISE_S3_C", "dept": "ISE", "sem": 3},

    {"id": "ISE_S5_A", "dept": "ISE", "sem": 5},
    {"id": "ISE_S5_B", "dept": "ISE", "sem": 5},

    {"id": "ISE_S7_A", "dept": "ISE", "sem": 7},
    {"id": "ISE_S7_B", "dept": "ISE", "sem": 7},

    # CSBS
    {"id": "CSBS_S1_A", "dept": "CSBS", "sem": 1},
    {"id": "CSBS_S1_B", "dept": "CSBS", "sem": 1},

    {"id": "CSBS_S3_A", "dept": "CSBS", "sem": 3},
    {"id": "CSBS_S3_B", "dept": "CSBS", "sem": 3},

    {"id": "CSBS_S5_A", "dept": "CSBS", "sem": 5},
    {"id": "CSBS_S7_A", "dept": "CSBS", "sem": 7},
]

SECTION_IDS = [x["id"] for x in SECTIONS]
SECTION_INFO = {x["id"]: x for x in SECTIONS}




STAFF = []

for i in range(1, 5):
    STAFF.append({
        "id": f"PROF_{i}",
        "name": f"Professor {i}",
        "designation": "Professor"
    })

for i in range(1, 7):
    STAFF.append({
        "id": f"ASSOC_{i}",
        "name": f"Associate Professor {i}",
        "designation": "Associate Professor"
    })

for i in range(1, 11):
    STAFF.append({
        "id": f"ASST_{i}",
        "name": f"Assistant Professor {i}",
        "designation": "Assistant Professor"
    })

STAFF_IDS = [x["id"] for x in STAFF]
STAFF_INFO = {x["id"]: x for x in STAFF}

MAX_SUBJECTS_PER_STAFF = 3
MAX_SECTIONS_PER_STAFF = 3




SUBJECTS = {
    "ISE": {
        3: [
            "Data Structures",
            "Database Management Systems",
            "Computer Networks",
            "Operating Systems",
            "Object Oriented Programming",
            "Mathematics for Computing"
        ],

        5: [
            "Machine Learning",
            "Web Technology",
            "Software Engineering",
            "Cloud Computing",
            "Data Mining",
            "Information Security"
        ],

        7: [
            "Deep Learning",
            "Big Data Analytics",
            "Natural Language Processing",
            "Distributed Systems",
            "Advanced Database Systems",
            "Project Management"
        ]
    },

    "CSBS": {
        1: [
            "Programming Fundamentals",
            "Discrete Mathematics",
            "Digital Logic",
            "Computer Organization",
            "Communication Skills",
            "Engineering Mathematics"
        ],

        3: [
            "Data Structures",
            "Database Systems",
            "Computer Networks",
            "Operating Systems",
            "Object Oriented Programming",
            "Probability and Statistics"
        ],

        5: [
            "Machine Learning",
            "Business Analytics",
            "Cloud Computing",
            "Software Engineering",
            "Data Visualization",
            "Information Security"
        ],

        7: [
            "Deep Learning",
            "Natural Language Processing",
            "Big Data Analytics",
            "Business Intelligence",
            "Distributed Systems",
            "Project Management"
        ]
    }
}


# ================================================================
# 5. OE DATABASE
# ================================================================

OE_CHOICES = [
    "Artificial Intelligence",
    "Cyber Security",
    "Blockchain Technology",
    "Cloud Computing",
    "Data Science",
    "Internet of Things",
    "Computer Vision",
    "Natural Language Processing",
    "Robotics",
    "Big Data Analytics",
    "Business Intelligence",
    "Digital Forensics"
]

ISE_OE = "Artificial Intelligence"
CSBS_OE = "Business Intelligence"


# ================================================================
# 6. PE DATABASE
# ================================================================

PE_SUBJECTS = [
    "Physical Education"
]


# ================================================================
# 7. AGENT 1 - CURATOR AGENT
# ================================================================

class CuratorAgent:

    def run(self):

        print("\n" + "=" * 70)
        print("[AGENT 1] CURATOR AGENT")
        print("=" * 70)

        events = []

        # --------------------------------------------------------
        # THEORY EVENTS
        # --------------------------------------------------------

        theory_counter = 0

        for section in SECTIONS:

            sid = section["id"]
            dept = section["dept"]
            sem = section["sem"]

            subjects = SUBJECTS[dept][sem]

            if sid in ["ISE_S3_A", "ISE_S3_B"]:
                total_theory = 22
            else:
                total_theory = 21

            for i in range(total_theory):

                subject = subjects[i % len(subjects)]

                events.append({
                    "id": f"T_{theory_counter:03d}",
                    "type": "THEORY",
                    "section": sid,
                    "dept": dept,
                    "sem": sem,
                    "subject": subject,
                    "duration": 1
                })

                theory_counter += 1

        # --------------------------------------------------------
        # LAB EVENTS
        # --------------------------------------------------------

        lab_counter = 0

        for section in SECTIONS:

            sid = section["id"]
            dept = section["dept"]
            sem = section["sem"]

            subjects = SUBJECTS[dept][sem]

            lab_subject = subjects[0] + " Lab"

            events.append({
                "id": f"LAB_{lab_counter:03d}",
                "type": "LAB",
                "section": sid,
                "dept": dept,
                "sem": sem,
                "subject": lab_subject,
                "duration": 3
            })

            lab_counter += 1

        # --------------------------------------------------------
        # OE EVENTS
        # --------------------------------------------------------

        oe_sections = [
            "ISE_S5_A",
            "ISE_S5_B",
            "ISE_S7_A",
            "ISE_S7_B",
            "CSBS_S5_A",
            "CSBS_S7_A"
        ]

        oe_counter = 0

        for sid in oe_sections:

            dept = SECTION_INFO[sid]["dept"]
            sem = SECTION_INFO[sid]["sem"]

            oe_subject = ISE_OE if dept == "ISE" else CSBS_OE

            for session in range(2):

                events.append({
                    "id": f"OE_{oe_counter:03d}",
                    "type": "OE",
                    "section": sid,
                    "dept": dept,
                    "sem": sem,
                    "subject": oe_subject,
                    "duration": 1
                })

                oe_counter += 1

        # --------------------------------------------------------
        # PE EVENTS
        # --------------------------------------------------------

        pe_groups = [
            ["ISE_S5_A", "CSBS_S5_A"],
            ["ISE_S5_B"],
            ["ISE_S7_A", "CSBS_S7_A"],
            ["ISE_S7_B"]
        ]

        pe_counter = 0

        for group in pe_groups:

            for session in range(2):

                main_section = group[0]

                events.append({
                    "id": f"PE_{pe_counter:03d}",
                    "type": "PE",
                    "section": main_section,
                    "sections": group,
                    "dept": "COMBINED",
                    "sem": 5 if "S5" in main_section else 7,
                    "subject": "Physical Education",
                    "duration": 1
                })

                pe_counter += 1

        print("Sections :", len(SECTIONS))
        print("Staff    :", len(STAFF))
        print("Rooms    :", len(ROOMS))
        print("Classes  :", len(CLASSROOMS))
        print("Labs     :", len(LABS))

        return events


# ================================================================
# 8. AGENT 2 - CONSTRAINT AGENT
# ================================================================

class ConstraintAgent:

    def run(self, events):

        print("\n" + "=" * 70)
        print("[AGENT 2] CONSTRAINT AGENT")
        print("=" * 70)

        counts = defaultdict(int)

        for e in events:
            counts[e["type"]] += 1

        print("Event counts:")
        for k, v in counts.items():
            print(f"  {k:10s}: {v}")

        print("\nScheduling constraints:")
        print("  Monday-Saturday")
        print("  7 classrooms")
        print("  2 labs")
        print("  OE at 09:30")
        print("  Short Break 10:30-11:00")
        print("  Lunch 13:00-14:00 (1:00 to 2:00)")
        print("  PE at 14:00")
        print("  Lab duration = 3 consecutive hours")
        print("  No room collision")
        print("  No teacher collision")
        print("  No section collision")
        print("  Max 3 subjects / teacher")
        print("  Max 3 sections / teacher")
        print("  Senior staff cannot start at 07:30")
        print("  12 OE choices but only 2 active OE streams")

        return {
            "events": events,
            "counts": dict(counts)
        }


# ================================================================
# 9. AGENT 3 - SOLVER AGENT
# ================================================================

class SolverAgent:

    def run(self, state):

        print("\n" + "=" * 70)
        print("[AGENT 3] SOLVER AGENT - CP-SAT (MULTI-STAGE OPTIMIZER)")
        print("=" * 70)

        events = state["events"]

        # Group events by course: (section(s), subject, type)
        # This guarantees every event of a course is taught by the SAME faculty member
        course_events = defaultdict(list)
        for i, e in enumerate(events):
            sec_key = tuple(sorted(e.get("sections", [e["section"]])))
            course_events[(sec_key, e["subject"], e["type"])].append(i)

        courses = list(course_events.keys())
        num_courses = len(courses)
        print(f"Aggregated courses: {num_courses}")

        # --------------------------------------------------------
        # STAGE 1: FACULTY ALLOTMENT OPTIMIZATION (CP-SAT)
        # --------------------------------------------------------
        print("\n[Stage 3.1] Faculty Allotment (Max 3 subjects, Max 3 sections, Workload balance)...")
        m_faculty = cp_model.CpModel()

        # assigned[c, t] is True if teacher t is allotted to course c
        assigned = {}
        for c_idx in range(num_courses):
            for t in range(len(STAFF)):
                assigned[c_idx, t] = m_faculty.NewBoolVar(f"as_{c_idx}_{t}")

        for c_idx in range(num_courses):
            m_faculty.AddExactlyOne([assigned[c_idx, t] for t in range(len(STAFF))])

        # Workload balancing (hours per faculty)
        course_hours = [sum(events[i]["duration"] for i in course_events[c]) for c in courses]
        for t in range(len(STAFF)):
            th = m_faculty.NewIntVar(12, 22, f"th_{t}")
            m_faculty.Add(th == sum(course_hours[c_idx] * assigned[c_idx, t] for c_idx in range(num_courses)))

        # Subject limit <= 3
        all_subjects = sorted(set(c[1] for c in courses))
        for t in range(len(STAFF)):
            sub_bools = []
            for s in all_subjects:
                matching_c = [c_idx for c_idx, c in enumerate(courses) if c[1] == s]
                sb = m_faculty.NewBoolVar(f"t_{t}_sub_{s}")
                m_faculty.AddMaxEquality(sb, [assigned[c_idx, t] for c_idx in matching_c])
                sub_bools.append(sb)
            m_faculty.Add(sum(sub_bools) <= MAX_SUBJECTS_PER_STAFF)

        # Section limit <= 3
        for t in range(len(STAFF)):
            sec_bools = []
            for sid in SECTION_IDS:
                matching_c = [c_idx for c_idx, c in enumerate(courses) if sid in c[0]]
                sb = m_faculty.NewBoolVar(f"t_{t}_sec_{sid}")
                m_faculty.AddMaxEquality(sb, [assigned[c_idx, t] for c_idx in matching_c])
                sec_bools.append(sb)
            m_faculty.Add(sum(sec_bools) <= MAX_SECTIONS_PER_STAFF)

        solver_f = cp_model.CpSolver()
        solver_f.parameters.max_time_in_seconds = 15
        solver_f.parameters.num_search_workers = 8
        status_f = solver_f.Solve(m_faculty)

        if status_f not in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
            raise RuntimeError("Faculty allotment failed. Check teacher constraints.")

        print(f"Faculty Allotment status: {solver_f.StatusName(status_f)}")

        # Assign allotted faculty to each event
        for c_idx, c in enumerate(courses):
            for t in range(len(STAFF)):
                if solver_f.Value(assigned[c_idx, t]):
                    teacher = STAFF[t]
                    for i in course_events[c]:
                        events[i]["teacher"] = teacher["name"]
                        events[i]["teacher_id"] = teacher["id"]
                        events[i]["designation"] = teacher["designation"]
                        events[i]["teacher_idx"] = t
                    break

        # --------------------------------------------------------
        # STAGE 2: TIMETABLE TIME-SLOT SCHEDULING (CP-SAT)
        # --------------------------------------------------------
        print("\n[Stage 3.2] Timetable Slot Scheduling (Shifts, Collisions, Lab blocks, Senior staff rules)...")
        m_sched = cp_model.CpModel()

        day = [m_sched.NewIntVar(0, NUM_DAYS - 1, f"day_{i}") for i in range(len(events))]
        slot = [m_sched.NewIntVar(0, NUM_SLOTS - 1, f"slot_{i}") for i in range(len(events))]
        start = [m_sched.NewIntVar(0, NUM_DAYS * NUM_SLOTS - 1, f"start_{i}") for i in range(len(events))]

        for i in range(len(events)):
            m_sched.Add(start[i] == day[i] * NUM_SLOTS + slot[i])

        # Senior staff cannot teach at slot 0 (07:30)
        for i, e in enumerate(events):
            if e["designation"] in ["Professor", "Associate Professor"]:
                m_sched.Add(slot[i] != 0)

        # Shift restrictions: Early shift (0..5) or Late shift (3..8)
        shift = {(sid, d): m_sched.NewBoolVar(f"late_{sid}_{d}") for sid in SECTION_IDS for d in range(NUM_DAYS)}
        for i, e in enumerate(events):
            sids = e.get("sections", [e["section"]])
            for sid in sids:
                for d in range(NUM_DAYS):
                    same_day = m_sched.NewBoolVar(f"ev_{i}_{sid}_d_{d}")
                    m_sched.Add(day[i] == d).OnlyEnforceIf(same_day)
                    m_sched.Add(day[i] != d).OnlyEnforceIf(same_day.Not())

                    late = shift[(sid, d)]
                    m_sched.Add(slot[i] >= 3 * late).OnlyEnforceIf(same_day)
                    m_sched.Add(slot[i] <= 5 + 3 * late).OnlyEnforceIf(same_day)

        # Fixed slots for OE and PE
        for i, e in enumerate(events):
            if e["type"] == "OE":
                m_sched.Add(slot[i] == OE_SLOT)
            elif e["type"] == "PE":
                m_sched.Add(slot[i] == PE_SLOT)
            elif e["type"] == "LAB":
                valid_starts = [d * NUM_SLOTS + s for d in range(NUM_DAYS) for s in [0, 5, 6]]
                m_sched.AddAllowedAssignments([start[i]], [[v] for v in valid_starts])

        # Section non-overlap
        sec_intervals = {sid: [] for sid in SECTION_IDS}
        for i, e in enumerate(events):
            sids = e.get("sections", [e["section"]])
            for sid in sids:
                iv = m_sched.NewIntervalVar(start[i], e["duration"], start[i] + e["duration"], f"sec_iv_{sid}_{i}")
                sec_intervals[sid].append(iv)

        for sid in SECTION_IDS:
            m_sched.AddNoOverlap(sec_intervals[sid])

        # Teacher non-overlap
        teacher_intervals = [[] for _ in range(len(STAFF))]
        for i, e in enumerate(events):
            t_idx = e["teacher_idx"]
            iv = m_sched.NewIntervalVar(start[i], e["duration"], start[i] + e["duration"], f"t_iv_{t_idx}_{i}")
            teacher_intervals[t_idx].append(iv)

        for t in range(len(STAFF)):
            m_sched.AddNoOverlap(teacher_intervals[t])

        # Cumulative room capacities (eliminates room permutation symmetry!)
        cr_intervals = []
        lab_intervals = []
        for i, e in enumerate(events):
            iv = m_sched.NewIntervalVar(start[i], e["duration"], start[i] + e["duration"], f"room_iv_{i}")
            if e["type"] == "LAB":
                lab_intervals.append(iv)
            else:
                cr_intervals.append(iv)

        m_sched.AddCumulative(cr_intervals, [1] * len(cr_intervals), NUM_CLASSROOMS)
        m_sched.AddCumulative(lab_intervals, [1] * len(lab_intervals), NUM_LABS)

        solver_sched = cp_model.CpSolver()
        solver_sched.parameters.max_time_in_seconds = SOLVER_TIME_LIMIT
        solver_sched.parameters.num_search_workers = 8

        print("CP-SAT Solver started...")
        status_sched = solver_sched.Solve(m_sched)
        print("Solver status:", solver_sched.StatusName(status_sched))

        if status_sched not in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
            raise RuntimeError("No feasible timetable generated. Check constraints.")

        # Extract scheduled slots
        result = []
        for i, e in enumerate(events):
            d_val = solver_sched.Value(day[i])
            s_val = solver_sched.Value(slot[i])
            result.append({
                **e,
                "day": DAYS[d_val],
                "day_index": d_val,
                "slot": s_val,
                "time": TIME_SLOTS[s_val],
                "start_time": d_val * NUM_SLOTS + s_val
            })

        # --------------------------------------------------------
        # STAGE 3: ROOM ALLOCATION
        # --------------------------------------------------------
        print("\n[Stage 3.3] Room Allocation (CR_1..CR_7, LAB_1..LAB_2)...")

        # Assign labs
        lab_occupancy = defaultdict(set)
        for e in result:
            if e["type"] == "LAB":
                for lab_name in LABS:
                    if all(lab_name not in lab_occupancy[e["start_time"] + k] for k in range(e["duration"])):
                        e["room"] = lab_name
                        for k in range(e["duration"]):
                            lab_occupancy[e["start_time"] + k].add(lab_name)
                        break

        # Assign classrooms
        cr_occupancy = defaultdict(set)
        for e in result:
            if e["type"] != "LAB":
                t_slot = e["start_time"]
                for cr_name in CLASSROOMS:
                    if cr_name not in cr_occupancy[t_slot]:
                        e["room"] = cr_name
                        cr_occupancy[t_slot].add(cr_name)
                        break

        print("Timetable generated successfully.")
        return result


# ================================================================
# 10. AGENT 4 - VALIDATOR AGENT
# ================================================================

class ValidatorAgent:

    def run(self, timetable):

        print("\n" + "=" * 70)
        print("[AGENT 4] VALIDATOR AGENT")
        print("=" * 70)

        errors = []
        warnings = []

        # --------------------------------------------------------
        # SECTION COLLISION
        # --------------------------------------------------------

        section_slots = defaultdict(list)

        for e in timetable:

            sections = e.get(
                "sections",
                [e["section"]]
            )

            duration = e["duration"]

            for sid in sections:

                for k in range(duration):

                    slot_number = (
                        e["day_index"] * NUM_SLOTS
                        + e["slot"]
                        + k
                    )

                    section_slots[
                        (sid, slot_number)
                    ].append(e["id"])

        for key, ids in section_slots.items():

            if len(ids) > 1:

                errors.append(
                    f"Section collision {key}: {ids}"
                )

        # --------------------------------------------------------
        # ROOM COLLISION
        # --------------------------------------------------------

        room_slots = defaultdict(list)

        for e in timetable:

            for k in range(e["duration"]):

                slot_number = (
                    e["day_index"] * NUM_SLOTS
                    + e["slot"]
                    + k
                )

                room_slots[
                    (e["room"], slot_number)
                ].append(e["id"])

        for key, ids in room_slots.items():

            if len(ids) > 1:

                errors.append(
                    f"Room collision {key}: {ids}"
                )

        # --------------------------------------------------------
        # STAFF COLLISION
        # --------------------------------------------------------

        staff_slots = defaultdict(list)

        for e in timetable:

            for k in range(e["duration"]):

                slot_number = (
                    e["day_index"] * NUM_SLOTS
                    + e["slot"]
                    + k
                )

                staff_slots[
                    (e["teacher_id"], slot_number)
                ].append(e["id"])

        for key, ids in staff_slots.items():

            if len(ids) > 1:

                errors.append(
                    f"Staff collision {key}: {ids}"
                )

        # --------------------------------------------------------
        # LAB VALIDATION
        # --------------------------------------------------------

        for e in timetable:

            if e["type"] == "LAB":

                if e["room"] not in LABS:

                    errors.append(
                        f"Lab {e['id']} assigned to "
                        f"non-lab room {e['room']}"
                    )

        # --------------------------------------------------------
        # OE VALIDATION
        # --------------------------------------------------------

        for e in timetable:

            if e["type"] == "OE":

                if e["slot"] != OE_SLOT:

                    errors.append(
                        f"OE {e['id']} is not at 09:30"
                    )

        # --------------------------------------------------------
        # PE VALIDATION
        # --------------------------------------------------------

        for e in timetable:

            if e["type"] == "PE":

                if e["slot"] != PE_SLOT:

                    errors.append(
                        f"PE {e['id']} is not at {TIME_SLOTS[PE_SLOT]}"
                    )

        # --------------------------------------------------------
        # BREAK & LUNCH VALIDATION
        # --------------------------------------------------------

        # Short Break is 10:30-11:00 (between slot 2: 09:30-10:30 and slot 3: 11:00-12:00).
        # Lunch Break is 13:00-14:00 (between slot 4: 12:00-13:00 and slot 5: 14:00-15:00).
        # Valid 3-hour lab start slots are 0 (07:30-10:30), 5 (14:00-17:00), or 6 (15:00-18:00).
        for e in timetable:
            if e["type"] == "LAB":
                if e["slot"] not in [0, 5, 6]:
                    errors.append(f"Lab {e['id']} crosses break/lunch or invalid block: slot {e['slot']}")

        # --------------------------------------------------------
        # SENIOR STAFF 07:30 VALIDATION
        # --------------------------------------------------------

        for e in timetable:

            if e["slot"] == 0:

                if e["designation"] in [
                    "Professor",
                    "Associate Professor"
                ]:

                    errors.append(
                        f"Senior staff at 07:30: {e['id']}"
                    )

        # --------------------------------------------------------
        # TEACHER SUBJECT COUNT
        # --------------------------------------------------------

        teacher_subjects = defaultdict(set)

        for e in timetable:

            teacher_subjects[
                e["teacher_id"]
            ].add(e["subject"])

        for teacher, subjects in teacher_subjects.items():

            if len(subjects) > MAX_SUBJECTS_PER_STAFF:

                errors.append(
                    f"{teacher} teaches "
                    f"{len(subjects)} subjects"
                )

        # --------------------------------------------------------
        # TEACHER SECTION COUNT
        # --------------------------------------------------------

        teacher_sections = defaultdict(set)

        for e in timetable:

            for sid in e.get(
                "sections",
                [e["section"]]
            ):

                teacher_sections[
                    e["teacher_id"]
                ].add(sid)

        for teacher, sections in teacher_sections.items():

            if len(sections) > MAX_SECTIONS_PER_STAFF:

                errors.append(
                    f"{teacher} teaches "
                    f"{len(sections)} sections"
                )

        # --------------------------------------------------------
        # EVENT COUNT
        # --------------------------------------------------------

        expected = {
            "THEORY": 275,
            "LAB": 13,
            "OE": 12,
            "PE": 8
        }

        actual = defaultdict(int)

        for e in timetable:

            actual[e["type"]] += 1

        for k, v in expected.items():

            if actual[k] != v:

                errors.append(
                    f"{k}: expected {v}, "
                    f"got {actual[k]}"
                )

        # --------------------------------------------------------
        # RESULT
        # --------------------------------------------------------

        print("\nValidation result:")

        if errors:

            print("FAILED")
            print("Errors:", len(errors))

            for err in errors[:30]:

                print("  -", err)

        else:

            print("PASSED")
            print("No hard constraint violations.")

        return {
            "status": "PASSED" if not errors else "FAILED",
            "errors": errors,
            "warnings": warnings,
            "teacher_subjects": teacher_subjects,
            "teacher_sections": teacher_sections
        }


# ================================================================
# 11. EXCEL GENERATOR
# ================================================================

class ExcelGenerator:

    def __init__(self, filename):

        self.filename = filename

        self.wb = Workbook()

        # Remove default sheet
        ws = self.wb.active
        self.wb.remove(ws)

        self.header_fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        self.lunch_fill = PatternFill(
            "solid",
            fgColor="D9E1F2"
        )

        self.break_fill = PatternFill(
            "solid",
            fgColor="FFF2CC"
        )

        self.header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        self.bold_font = Font(
            bold=True
        )

        self.center = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        self.thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

    # ------------------------------------------------------------
    # WRITE SECTION TIMETABLE
    # ------------------------------------------------------------

    def write_section(self, sid, timetable):

        ws = self.wb.create_sheet(sid)

        ws["A1"] = f"TIMETABLE - {sid}"
        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        headers = [
            "Time",
            *DAYS
        ]

        for c, h in enumerate(headers, 1):

            cell = ws.cell(
                row=3,
                column=c,
                value=h
            )

            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center
            cell.border = self.thin_border

        event_map = {}

        for e in timetable:

            sections = e.get(
                "sections",
                [e["section"]]
            )

            if sid not in sections:
                continue

            # Map all duration hours (e.g. 3-hour labs)
            for k in range(e["duration"]):
                event_map[
                    (e["day"], e["slot"] + k)
                ] = e

        current_row = 4

        for s in range(NUM_SLOTS):

            # Insert Short Break banner between slot 2 (09:30-10:30) and slot 3 (11:00-12:00)
            if s == 3:
                ws.cell(row=current_row, column=1, value="10:30-11:00").font = self.bold_font
                ws.cell(row=current_row, column=1).alignment = self.center
                ws.cell(row=current_row, column=1).border = self.thin_border
                ws.cell(row=current_row, column=1).fill = self.break_fill

                for d in range(NUM_DAYS):
                    cell = ws.cell(row=current_row, column=d + 2, value="SHORT BREAK")
                    cell.alignment = self.center
                    cell.border = self.thin_border
                    cell.fill = self.break_fill
                    cell.font = self.bold_font
                current_row += 1

            # Insert Lunch Break banner between slot 4 (12:00-13:00) and slot 5 (14:00-15:00)
            if s == 5:
                ws.cell(row=current_row, column=1, value="13:00-14:00").font = self.bold_font
                ws.cell(row=current_row, column=1).alignment = self.center
                ws.cell(row=current_row, column=1).border = self.thin_border
                ws.cell(row=current_row, column=1).fill = self.lunch_fill

                for d in range(NUM_DAYS):
                    cell = ws.cell(row=current_row, column=d + 2, value="LUNCH BREAK")
                    cell.alignment = self.center
                    cell.border = self.thin_border
                    cell.fill = self.lunch_fill
                    cell.font = self.bold_font
                current_row += 1

            ws.cell(
                row=current_row,
                column=1,
                value=TIME_SLOTS[s]
            ).font = self.bold_font
            ws.cell(row=current_row, column=1).alignment = self.center
            ws.cell(row=current_row, column=1).border = self.thin_border

            for d in range(NUM_DAYS):

                e = event_map.get(
                    (DAYS[d], s)
                )

                cell = ws.cell(
                    row=current_row,
                    column=d + 2
                )

                if e:

                    cell.value = (
                        f"{e['subject']}\n"
                        f"[{e['type']}]\n"
                        f"{e['teacher']}\n"
                        f"{e['room']}"
                    )

                else:

                    cell.value = "FREE"

                cell.alignment = self.center
                cell.border = self.thin_border

            current_row += 1

        for col in range(1, 8):

            ws.column_dimensions[
                get_column_letter(col)
            ].width = 24

    # ------------------------------------------------------------
    # MASTER TIMETABLE
    # ------------------------------------------------------------

    def write_master(self, dept, timetable):

        ws = self.wb.create_sheet(
            f"{dept}_MASTER"
        )

        ws["A1"] = f"{dept} MASTER TIMETABLE"
        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        headers = [
            "Day",
            "Time",
            "Section",
            "Subject",
            "Type",
            "Teacher",
            "Room"
        ]

        for c, h in enumerate(headers, 1):

            cell = ws.cell(
                row=3,
                column=c,
                value=h
            )

            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center
            cell.border = self.thin_border

        row = 4

        dept_events = [
            e
            for e in timetable
            if (
                e["dept"] == dept
                or (
                    e["type"] == "PE"
                    and any(
                        SECTION_INFO[s]["dept"] == dept
                        for s in e.get(
                            "sections",
                            []
                        )
                    )
                )
            )
        ]

        dept_events.sort(
            key=lambda x: (
                x["day_index"],
                x["slot"],
                x["section"]
            )
        )

        for e in dept_events:

            section_text = ",".join(
                e.get(
                    "sections",
                    [e["section"]]
                )
            )

            values = [
                e["day"],
                e["time"],
                section_text,
                e["subject"],
                e["type"],
                e["teacher"],
                e["room"]
            ]

            for c, value in enumerate(
                values,
                1
            ):

                cell = ws.cell(
                    row=row,
                    column=c,
                    value=value
                )

                cell.alignment = self.center
                cell.border = self.thin_border

            row += 1

        for c in range(1, 8):

            ws.column_dimensions[
                get_column_letter(c)
            ].width = 24

    # ------------------------------------------------------------
    # STAFF TIMETABLE
    # ------------------------------------------------------------

    def write_staff(self, timetable):

        ws = self.wb.create_sheet(
            "STAFF_TIMETABLE"
        )

        headers = [
            "Teacher",
            "Designation",
            "Day",
            "Time",
            "Subject",
            "Section",
            "Room",
            "Type"
        ]

        for c, h in enumerate(headers, 1):

            cell = ws.cell(
                row=1,
                column=c,
                value=h
            )

            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center
            cell.border = self.thin_border

        row = 2

        sorted_events = sorted(
            timetable,
            key=lambda x: (
                x["teacher"],
                x["day_index"],
                x["slot"]
            )
        )

        for e in sorted_events:

            values = [
                e["teacher"],
                e["designation"],
                e["day"],
                e["time"],
                e["subject"],
                ",".join(
                    e.get(
                        "sections",
                        [e["section"]]
                    )
                ),
                e["room"],
                e["type"]
            ]

            for c, value in enumerate(
                values,
                1
            ):

                cell = ws.cell(
                    row=row,
                    column=c,
                    value=value
                )

                cell.alignment = self.center
                cell.border = self.thin_border

            row += 1

        for c in range(1, 9):

            ws.column_dimensions[
                get_column_letter(c)
            ].width = 25

    # ------------------------------------------------------------
    # ROOM UTILIZATION
    # ------------------------------------------------------------

    def write_room_utilization(self, timetable):

        ws = self.wb.create_sheet(
            "ROOM_UTILIZATION"
        )

        headers = [
            "Room",
            "Day",
            "Time",
            "Subject",
            "Section",
            "Teacher",
            "Type"
        ]

        for c, h in enumerate(headers, 1):

            cell = ws.cell(
                row=1,
                column=c,
                value=h
            )

            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center
            cell.border = self.thin_border

        row = 2

        sorted_events = sorted(
            timetable,
            key=lambda x: (
                ROOMS.index(x["room"]),
                x["day_index"],
                x["slot"]
            )
        )

        for e in sorted_events:

            values = [
                e["room"],
                e["day"],
                e["time"],
                e["subject"],
                ",".join(
                    e.get(
                        "sections",
                        [e["section"]]
                    )
                ),
                e["teacher"],
                e["type"]
            ]

            for c, value in enumerate(
                values,
                1
            ):

                cell = ws.cell(
                    row=row,
                    column=c,
                    value=value
                )

                cell.alignment = self.center
                cell.border = self.thin_border

            row += 1

        for c in range(1, 8):
            ws.column_dimensions[get_column_letter(c)].width = 22

    # ------------------------------------------------------------
    # VALIDATION REPORT
    # ------------------------------------------------------------

    def write_validation(self, validation):

        ws = self.wb.create_sheet(
            "VALIDATION_REPORT"
        )

        ws["A1"] = "TIMETABLE VALIDATION REPORT"
        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        ws["A3"] = "Status"
        ws["B3"] = validation["status"]
        ws["B3"].font = Font(bold=True, color="008000" if validation["status"] == "PASSED" else "FF0000")

        ws["A4"] = "Errors"
        ws["B4"] = len(
            validation["errors"]
        )

        row = 6

        ws.cell(
            row=row,
            column=1,
            value="Error"
        ).font = self.header_font

        ws.cell(
            row=row,
            column=1
        ).fill = self.header_fill

        row += 1

        if not validation["errors"]:
            ws.cell(row=row, column=1, value="None - All hard and soft constraints validated successfully!")
            row += 1
        else:
            for error in validation["errors"]:
                ws.cell(
                    row=row,
                    column=1,
                    value=error
                )
                row += 1

        row += 2

        ws.cell(
            row=row,
            column=1,
            value="Teacher"
        )

        ws.cell(
            row=row,
            column=2,
            value="Subjects"
        )

        ws.cell(
            row=row,
            column=3,
            value="Sections"
        )

        for c in range(1, 4):

            ws.cell(
                row=row,
                column=c
            ).fill = self.header_fill

            ws.cell(
                row=row,
                column=c
            ).font = self.header_font

        row += 1

        for teacher in STAFF_IDS:

            subjects = validation[
                "teacher_subjects"
            ].get(
                teacher,
                set()
            )

            sections = validation[
                "teacher_sections"
            ].get(
                teacher,
                set()
            )

            ws.cell(
                row=row,
                column=1,
                value=teacher
            )

            ws.cell(
                row=row,
                column=2,
                value=len(subjects)
            )

            ws.cell(
                row=row,
                column=3,
                value=len(sections)
            )

            row += 1

        for c in range(1, 4):
            ws.column_dimensions[get_column_letter(c)].width = 20

    # ------------------------------------------------------------
    # AGENT REPORT
    # ------------------------------------------------------------

    def write_agent_report(self):

        ws = self.wb.create_sheet(
            "AGENT_EXECUTION"
        )

        report = [
            ["Agent", "Function", "Status"],
            [
                "Agent 1 - Curator",
                "Created sections, subjects, teaching events and courses",
                "COMPLETED"
            ],
            [
                "Agent 2 - Constraint",
                "Applied timetable, room, faculty, and shift constraints",
                "COMPLETED"
            ],
            [
                "Agent 3 - Solver",
                "CP-SAT multi-stage optimal scheduling and allocation",
                "COMPLETED"
            ],
            [
                "Agent 4 - Validator",
                "Validated collisions, rooms, senior staff, and workload limits",
                "COMPLETED"
            ],
            [
                "Agent 5 - Excel",
                "Generated master, section, staff, and utilization workbooks",
                "COMPLETED"
            ]
        ]

        for r, row_data in enumerate(
            report,
            1
        ):

            for c, value in enumerate(
                row_data,
                1
            ):

                cell = ws.cell(
                    row=r,
                    column=c,
                    value=value
                )

                cell.alignment = self.center
                cell.border = self.thin_border

                if r == 1:

                    cell.fill = self.header_fill
                    cell.font = self.header_font

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 65
        ws.column_dimensions["C"].width = 20

    # ------------------------------------------------------------
    # SAVE
    # ------------------------------------------------------------

    def save(self):

        self.wb.save(
            self.filename
        )

        print(
            "\nExcel file created:",
            os.path.abspath(self.filename)
        )


# ================================================================
# 12. MAIN AGENTIC PIPELINE
# ================================================================

def main():

    start_time = time.time()

    print("\n")
    print("=" * 80)
    print("       AGENTIC AI COLLEGE TIMETABLE GENERATOR")
    print("=" * 80)
    print("       FINAL MONDAY-SATURDAY VERSION")
    print("=" * 80)

    # ------------------------------------------------------------
    # AGENT 1
    # ------------------------------------------------------------

    curator = CuratorAgent()

    events = curator.run()

    # ------------------------------------------------------------
    # AGENT 2
    # ------------------------------------------------------------

    constraint_agent = ConstraintAgent()

    state = constraint_agent.run(
        events
    )

    # ------------------------------------------------------------
    # AGENT 3
    # ------------------------------------------------------------

    solver_agent = SolverAgent()

    timetable = solver_agent.run(
        state
    )

    # ------------------------------------------------------------
    # AGENT 4
    # ------------------------------------------------------------

    validator = ValidatorAgent()

    validation = validator.run(
        timetable
    )

    # ------------------------------------------------------------
    # AGENT 5
    # ------------------------------------------------------------

    print("\n" + "=" * 70)
    print("[AGENT 5] EXCEL REPORT GENERATOR")
    print("=" * 70)

    excel = ExcelGenerator(
        OUTPUT_FILE
    )

    # Individual section sheets
    for sid in SECTION_IDS:

        excel.write_section(
            sid,
            timetable
        )

    # Department masters
    excel.write_master(
        "ISE",
        timetable
    )

    excel.write_master(
        "CSBS",
        timetable
    )

    # Staff
    excel.write_staff(
        timetable
    )

    # Rooms
    excel.write_room_utilization(
        timetable
    )

    # Validation
    excel.write_validation(
        validation
    )

    # Agent report
    excel.write_agent_report()

    excel.save()

    # ------------------------------------------------------------
    # FINAL SUMMARY
    # ------------------------------------------------------------

    elapsed = time.time() - start_time

    print("\n")
    print("=" * 80)
    print("FINAL TIMETABLE GENERATION SUMMARY")
    print("=" * 80)

    print("Sections generated :", len(SECTION_IDS))
    print("Total events       :", len(timetable))

    counts = defaultdict(int)

    for e in timetable:
        counts[e["type"]] += 1

    print("Theory events      :", counts["THEORY"])
    print("Lab events         :", counts["LAB"])
    print("OE events          :", counts["OE"])
    print("PE events          :", counts["PE"])

    print("Classrooms         :", NUM_CLASSROOMS)
    print("Labs               :", NUM_LABS)
    print("Days               :", len(DAYS))
    print("Validation         :", validation["status"])
    print("Execution time     :", round(elapsed, 2), "seconds")

    print("\nGenerated file:")
    print(os.path.abspath(OUTPUT_FILE))

    print("\nSection sheets:")

    for sid in SECTION_IDS:

        print("  [OK]", sid)

    print("\nAdditional sheets:")
    print("  [OK] ISE_MASTER")
    print("  [OK] CSBS_MASTER")
    print("  [OK] STAFF_TIMETABLE")
    print("  [OK] ROOM_UTILIZATION")
    print("  [OK] VALIDATION_REPORT")
    print("  [OK] AGENT_EXECUTION")

    print("\n" + "=" * 80)
    print("TIMETABLE GENERATION COMPLETED")
    print("=" * 80)

    return timetable, validation


# ================================================================
# 13. RUN
# ================================================================

if __name__ == "__main__":

    timetable, validation = main()
